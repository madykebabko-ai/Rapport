"""
Automatisation Excel Bancaire — Groq API
=========================================
Lit le PDF des états financiers tunisiens et remplit automatiquement
la colonne 31/12/2024 du modèle Excel bancaire.

Unité dans le modèle Excel : KTND (milliers de dinars tunisiens)
Unité dans le PDF            : DT  (dinars tunisiens)
→ Le script divise toutes les valeurs PDF par 1000 avant de les écrire

Installation : pip install groq pdfplumber openpyxl python-dotenv
Clé API gratuite : https://console.groq.com
"""

import json
import os
from datetime import datetime
import pdfplumber
import openpyxl
from groq import Groq
from dotenv import load_dotenv

# Charge les variables d'environnement depuis .env
load_dotenv()

# ─────────────────────────────────────────────
# CONFIGURATION — modifier ces 3 lignes
# ─────────────────────────────────────────────

API_KEY      = os.getenv("GROQ_API_KEY")
CHEMIN_PDF   = "Sotipapier EF indiv 31.12.2024.pdf"
CHEMIN_EXCEL = "Modele décompo.xlsx"

DATE_CIBLE_NUM      = 45657                    # Numéro de série Excel pour 31/12/2024
DATE_CIBLE_DATETIME = datetime(2024, 12, 31)   # Format datetime pour 31/12/2024

if not API_KEY:
    raise ValueError("❌ Clé API manquante. Crée un fichier .env avec GROQ_API_KEY=ta_clé")


# ─────────────────────────────────────────────
# ÉTAPE 1 : Extraction du texte PDF (30 pages max)
# ─────────────────────────────────────────────

def extraire_texte_pdf(chemin_pdf: str) -> str:
    print(f"📄 Lecture du PDF : {chemin_pdf}")
    texte_total = []

    with pdfplumber.open(chemin_pdf) as pdf:
        for i, page in enumerate(pdf.pages[:30]):
            texte = page.extract_text()
            if texte:
                texte_total.append(f"--- PAGE {i+1} ---\n{texte}")

    if not texte_total:
        raise ValueError("❌ Impossible d'extraire le texte du PDF.")

    resultat = "\n\n".join(texte_total)

    if len(resultat) > 20000:
        resultat = resultat[:20000]
        print("⚠️  Texte tronqué à 20 000 caractères")

    print(f"✅ {len(texte_total)} pages extraites ({len(resultat)} caractères)")
    return resultat


# ─────────────────────────────────────────────
# ÉTAPE 2 : Extraction des données via Groq API
# ─────────────────────────────────────────────

def extraire_donnees_avec_groq(texte_pdf: str, api_key: str) -> dict:
    print("🤖 Envoi à Groq API (Llama 3)...")

    client = Groq(api_key=api_key)

    prompt = f"""Tu es un expert-comptable spécialisé dans les états financiers tunisiens.

Voici le contenu d'un PDF d'états financiers arrêtés au 31/12/2024, exprimé en Dinars Tunisiens (DT) :

{texte_pdf}

Extrait UNIQUEMENT les valeurs au 31/12/2024 (pas 2023).
Retourne UNIQUEMENT un JSON valide, sans texte avant ni après, sans balises markdown.
Les montants doivent être en DT entiers (pas en milliers). Si une valeur est introuvable, mets null.

{{
  "entreprise": "nom de la société",
  "immo_incorporelles_brutes": null,
  "amort_immo_incorporelles": null,
  "immo_corporelles_brutes": null,
  "amort_immo_corporelles": null,
  "immo_financieres_brutes": null,
  "prov_immo_financieres": null,
  "stocks_bruts": null,
  "prov_stocks": null,
  "clients_bruts": null,
  "prov_clients": null,
  "autres_actifs_courants": null,
  "placements_actifs_financiers": null,
  "liquidites": null,
  "capital_social": null,
  "prime_emission": null,
  "reserves_legales": null,
  "resultats_reportes": null,
  "resultat_net": null,
  "prov_risques_charges": null,
  "emprunts_long_terme": null,
  "fournisseurs": null,
  "autres_passifs_courants": null,
  "concours_bancaires": null,
  "chiffre_affaires": null,
  "charges_personnel": null,
  "autres_charges_exploitation": null,
  "charges_financieres_nettes": null,
  "impot_benefices": null
}}"""

    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {
                "role": "system",
                "content": "Tu es un expert-comptable. Tu réponds UNIQUEMENT en JSON valide, sans texte ni balises."
            },
            {"role": "user", "content": prompt}
        ],
        temperature=0,
        max_tokens=2048,
    )

    texte_reponse = response.choices[0].message.content.strip()
    print("✅ Réponse reçue de Groq")

    for tag in ["```json", "```"]:
        if texte_reponse.startswith(tag):
            texte_reponse = texte_reponse[len(tag):]
    if texte_reponse.endswith("```"):
        texte_reponse = texte_reponse[:-3]
    texte_reponse = texte_reponse.strip()

    donnees = json.loads(texte_reponse)
    print(f"✅ Données extraites pour : {donnees.get('entreprise', '?')}")
    return donnees


# ─────────────────────────────────────────────
# ÉTAPE 3 : Remplissage Excel — colonne 31/12/2024
# ─────────────────────────────────────────────

def remplir_excel(donnees: dict, chemin_excel: str) -> str:
    print(f"📊 Ouverture du fichier Excel : {chemin_excel}")

    wb = openpyxl.load_workbook(chemin_excel)

    if "Saisie" in wb.sheetnames:
        ws = wb["Saisie"]
        print("✅ Feuille 'Saisie' trouvée")
    else:
        ws = wb.active
        print(f"⚠️  Feuille 'Saisie' introuvable — utilisation de '{ws.title}'")

    # ── Trouver la colonne 31/12/2024
    col_cible = None
    ligne_entetes = None

    for row in ws.iter_rows():
        for cell in row:
            val = cell.value
            if isinstance(val, datetime):
                if val.year == 2024 and val.month == 12 and val.day == 31:
                    col_cible = cell.column
                    ligne_entetes = cell.row
                    break
            elif val == DATE_CIBLE_NUM:
                col_cible = cell.column
                ligne_entetes = cell.row
                break
            elif isinstance(val, str) and ("12/24" in val or "31/12/2024" in val):
                col_cible = cell.column
                ligne_entetes = cell.row
                break
        if col_cible:
            break

    if not col_cible:
        raise ValueError(
            "❌ La colonne 31/12/2024 est introuvable.\n"
            "   Vérifie que la date 31/12/2024 est bien présente dans la feuille Saisie."
        )

    print(f"✅ Colonne 31/12/2024 trouvée : colonne {col_cible}")

    # ── Mapping : code colonne A → clé JSON
    MAPPING = [
        ("AD+AF+AH+AJ+AL",       "immo_incorporelles_brutes"),
        ("AC+AE+AG+AI+AK+AM",    "amort_immo_incorporelles"),
        ("AN+AP+AR+AT+AV+AX",    "immo_corporelles_brutes"),
        ("AO+AQ+AS+AU+AW+AY",    "amort_immo_corporelles"),
        ("CS+CU",                "immo_financieres_brutes"),
        ("CT+CV",                "prov_immo_financieres"),
        ("BL",                   "stocks_bruts"),
        ("BM",                   "prov_stocks"),
        ("BX",                   "clients_bruts"),
        ("BY",                   "prov_clients"),
        ("BV-BW",                "autres_actifs_courants"),
        ("CD",                   "placements_actifs_financiers"),
        ("CF",                   "liquidites"),
        ("DA",                   "capital_social"),
        ("DB",                   "prime_emission"),
        ("DC+DD+DE+DF+DG",       "reserves_legales"),
        ("DH",                   "resultats_reportes"),
        ("DI",                   "resultat_net"),
        ("dans DR : part stable", "prov_risques_charges"),
        ("DT",                   "emprunts_long_terme"),
        ("DX",                   "fournisseurs"),
        ("DY",                   "autres_passifs_courants"),
        ("EH",                   "concours_bancaires"),
        ("FL",                   "chiffre_affaires"),
        ("FY",                   "charges_personnel"),
        ("FW",                   "autres_charges_exploitation"),
        ("GR",                   "charges_financieres_nettes"),
        ("HK",                   "impot_benefices"),
    ]

    # Index des codes présents dans la colonne A
    index_codes = {}
    for row in ws.iter_rows(min_row=ligne_entetes + 1):
        val_a = row[0].value
        if val_a is not None:
            index_codes[str(val_a).strip()] = row[0].row

    # ── Remplir les cellules (DT ÷ 1000 = KTND)
    cellules_remplies = 0
    cellules_manquantes = []

    for code, cle_json in MAPPING:
        valeur_dt = donnees.get(cle_json)

        if valeur_dt is None:
            cellules_manquantes.append(f"{code} ({cle_json})")
            continue

        if code not in index_codes:
            cellules_manquantes.append(f"{code} — code introuvable dans Excel")
            continue

        num_ligne = index_codes[code]
        valeur_ktnd = round(valeur_dt / 1000)
        ws.cell(row=num_ligne, column=col_cible, value=valeur_ktnd)
        print(f"   ✍️  {code:<30} → {valeur_ktnd:>10} KTND")
        cellules_remplies += 1

    # ── Sauvegarde
    entreprise = donnees.get("entreprise", "inconnu").replace(" ", "_")
    nom_sortie = f"REMPLI_{entreprise}_31122024.xlsx"
    wb.save(nom_sortie)

    print(f"\n✅ {cellules_remplies} cellules remplies")
    if cellules_manquantes:
        print(f"⚠️  {len(cellules_manquantes)} valeurs non trouvées :")
        for m in cellules_manquantes:
            print(f"   - {m}")
    print(f"✅ Fichier sauvegardé : {nom_sortie}")
    return nom_sortie


# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────

def main():
    print("=" * 60)
    print("  AUTOMATISATION EXCEL BANCAIRE — Groq API (Gratuit)")
    print("=" * 60)

    texte_pdf     = extraire_texte_pdf(CHEMIN_PDF)
    donnees       = extraire_donnees_avec_groq(texte_pdf, API_KEY)

    print("\n📋 JSON extrait par l'IA (en DT) :")
    print(json.dumps(donnees, ensure_ascii=False, indent=2))

    fichier_final = remplir_excel(donnees, CHEMIN_EXCEL)

    print("\n" + "=" * 60)
    print(f"  🎉 TERMINÉ ! → {fichier_final}")
    print("=" * 60)


if __name__ == "__main__":
    main()
